import os
import re
import openpyxl
import pandas as pd
import datetime

EXCEL_PATH = r"C:\Users\Admin\Downloads\BT.xlsx"
OUTPUT_DIR = r"e:\GRP-AC"

def clean_header(val):
    if val is None:
        return ""
    s = str(val).strip()
    # Replace non-ASCII characters (corrupted degree symbols) with standard degree symbol °
    s = ''.join([c if ord(c) < 128 else '°' for c in s])
    s = s.replace('°°', '°').strip()
    s = re.sub(r'\s+', ' ', s)
    return s

def format_cell_value(val):
    if val is None:
        return None
    if isinstance(val, datetime.time):
        return val.strftime("%H:%M:%S")
    if isinstance(val, datetime.datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    return val

def extract_table(ws, start_col, end_col, header_row=9, data_start_row=12):
    raw_headers = []
    col_indices = []
    for c in range(start_col, end_col + 1):
        h = clean_header(ws.cell(header_row, c).value)
        if not h:
            h = clean_header(ws.cell(header_row - 1, c).value)
        if not h:
            h = f"Col_{c}"
        raw_headers.append(h)
        col_indices.append(c)

    seen = {}
    headers = []
    for h in raw_headers:
        if h in seen:
            seen[h] += 1
            headers.append(f"{h}_{seen[h]}")
        else:
            seen[h] = 0
            headers.append(h)

    rows = []
    for r in range(data_start_row, ws.max_row + 1):
        row_vals = [format_cell_value(ws.cell(r, c).value) for c in col_indices]
        if any(v is not None for v in row_vals):
            rows.append(row_vals)

    df = pd.DataFrame(rows, columns=headers)
    return df

def main():
    print(f"Loading Excel file: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws_pi = wb["BT-AC10 Production Insights"]

    print("Extracting Table 1: Production Data (Cols 1-71)...")
    df_prod = extract_table(ws_pi, start_col=1, end_col=71, header_row=9, data_start_row=12)
    prod_path = os.path.join(OUTPUT_DIR, "production_data.csv")
    df_prod.to_csv(prod_path, index=False, encoding="utf-8-sig")

    print("Extracting Table 2: Orange Quality Data (Cols 73-84)...")
    df_orange_q = extract_table(ws_pi, start_col=73, end_col=84, header_row=9, data_start_row=12)
    orange_path = os.path.join(OUTPUT_DIR, "orange_quality_data.csv")
    df_orange_q.to_csv(orange_path, index=False, encoding="utf-8-sig")

    print("Extracting Table 3: Quality Evaluation Data (Cols 86-98)...")
    df_eval_q = extract_table(ws_pi, start_col=86, end_col=98, header_row=9, data_start_row=12)
    eval_path = os.path.join(OUTPUT_DIR, "quality_evaluation_data.csv")
    df_eval_q.to_csv(eval_path, index=False, encoding="utf-8-sig")

    ws_rep = wb["Report"]
    print("Extracting Raw Report Data (Sheet 'Report')...")
    rep_headers = [clean_header(ws_rep.cell(18, c).value) or f"Col_{c}" for c in range(1, ws_rep.max_column + 1)]
    rep_rows = []
    for r in range(19, ws_rep.max_row + 1):
        r_vals = [format_cell_value(ws_rep.cell(r, c).value) for c in range(1, ws_rep.max_column + 1)]
        if any(v is not None for v in r_vals):
            rep_rows.append(r_vals)
    df_rep = pd.DataFrame(rep_rows, columns=rep_headers)
    rep_path = os.path.join(OUTPUT_DIR, "raw_report_data.csv")
    df_rep.to_csv(rep_path, index=False, encoding="utf-8-sig")

    print("\n--- Extracted CSV Headers ---")
    print("Orange Quality Data Headers:", df_orange_q.columns.tolist())
    print("Production Data Headers Sample:", df_prod.columns[10:20].tolist())

if __name__ == "__main__":
    main()
